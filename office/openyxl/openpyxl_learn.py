from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

# ===== 写 Excel =====
wb = Workbook()
ws = wb.active
ws.title = "销售数据"

# 写表头（加粗 + 背景色）
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

for col, title in enumerate(["产品", "销量", "单价"], start=1):
    cell = ws.cell(row=1, column=col, value=title)
    cell.font = header_font
    cell.fill = header_fill

# 写数据
data = [("键盘", 120, 199), ("鼠标", 200, 89), ("显示器", 45, 1299)]
for row, item in enumerate(data, start=2):
    for col, value in enumerate(item, start=1):
        ws.cell(row=row, column=col, value=value)

wb.save("sales.xlsx")

# ===== 读 Excel =====
wb = load_workbook("sales.xlsx", data_only=True)  # data_only 取计算后的值
ws = wb["销售数据"]
for row in ws.iter_rows(min_row=2, values_only=True):
    print(row)          # ('键盘', 120, 199) ...
